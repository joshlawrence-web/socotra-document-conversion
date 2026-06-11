#!/usr/bin/env python3
"""
Leg 0 — Document Ingestion (PDF / Word → raw HTML) + Field/Conditional Extraction

Converts a customer's source document (PDF or Word) into a rough HTML file
suitable for the existing Leg 1 pipeline, then extracts {field_name} tokens
and [[conditional]] blocks, annotates them, and writes all pipeline-ready artifacts.

Usage:
    python3 scripts/leg0_ingest.py --input <path.docx|path.pdf> [--output-dir <dir>]
    python3 scripts/leg0_ingest.py --parse-conditional-form <filled-form.md> --output-dir <dir>

Outputs (normal mode):
    {stem}.raw.html           — raw converted HTML (pre-annotation)
    {stem}.annotated.html     — HTML with {field} → $TBD_field, [[cond]] → $doc.condN
    {stem}.mapping.yaml       — leg2-compatible mapping (pipeline input; enriched in-place by Leg 2)
    {stem}.conditional-form.md — customer-facing conditional form

Output (--parse-conditional-form mode):
    {stem}.conditional-registry.yaml — parsed conditional registry for Leg 4

Exit codes:
    0  success
    1  any error (file not found, unsupported format, import failure)
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import yaml

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _derive_accessor(velocity: str, category: str) -> str:
    """Derive clean accessor from velocity path + category."""
    v = (velocity or "").strip()
    cat = (category or "").strip()
    if cat == "system":
        return "policy." + v[len("$data."):] if v.startswith("$data.") else v
    if cat == "quote_system":
        return "quote." + v[len("$data."):] if v.startswith("$data.") else v
    if cat == "policy_data":
        return "policy." + v[len("$data."):] if v.startswith("$data.") else v
    if v.startswith("$data."):
        return v[len("$data."):]
    if v.startswith("$"):
        return v[1:]
    return v


def _escape(text: str) -> str:
    return (
        text.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
    )


# ---------------------------------------------------------------------------
# Word (.docx) → HTML
# ---------------------------------------------------------------------------

def convert_docx(docx_path: Path) -> str:
    try:
        from docx import Document
    except ImportError:
        sys.exit(
            "Missing dependency python-docx.\n"
            "Install with: pip install python-docx --break-system-packages"
        )

    doc = Document(str(docx_path))
    parts: list[str] = []

    def _paragraph_html(para) -> str | None:
        text = para.text.strip()
        if not text:
            return None
        style_name = (para.style.name or "").lower()
        if "heading 1" in style_name:
            return f"<h1>{_escape(text)}</h1>"
        if "heading 2" in style_name or "heading 3" in style_name:
            return f"<h2>{_escape(text)}</h2>"
        if "heading" in style_name:
            return f"<h2>{_escape(text)}</h2>"
        runs_with_text = [r for r in para.runs if r.text.strip()]
        if runs_with_text and all(r.bold for r in runs_with_text):
            return f"<p><strong>{_escape(text)}</strong></p>"
        return f"<p>{_escape(text)}</p>"

    def _table_html(table) -> str:
        rows = ["<table>"]
        for row in table.rows:
            cells = "".join(
                f"<td>{_escape(cell.text.strip())}</td>" for cell in row.cells
            )
            rows.append(f"  <tr>{cells}</tr>")
        rows.append("</table>")
        return "\n".join(rows)

    from docx.oxml.ns import qn  # noqa: F401,PLC0415

    for child in doc.element.body:
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag

        if tag == "p":
            from docx.text.paragraph import Paragraph  # noqa: PLC0415
            para = Paragraph(child, doc)
            html = _paragraph_html(para)
            if html:
                parts.append(html)

        elif tag == "tbl":
            from docx.table import Table  # noqa: PLC0415
            table = Table(child, doc)
            parts.append(_table_html(table))

    body_content = "\n".join(parts)
    return f"<html>\n<body>\n{body_content}\n</body>\n</html>\n"


# ---------------------------------------------------------------------------
# PDF → HTML
# ---------------------------------------------------------------------------

def convert_pdf(pdf_path: Path) -> str:
    try:
        import pdfplumber
    except ImportError:
        sys.exit(
            "Missing dependency pdfplumber.\n"
            "Install with: pip install pdfplumber --break-system-packages"
        )

    parts: list[str] = []

    with pdfplumber.open(str(pdf_path)) as pdf:
        for page_num, page in enumerate(pdf.pages):
            if page_num > 0:
                parts.append("<!-- page break -->")

            table_bboxes: list[tuple] = []
            try:
                tables = page.extract_tables()
                for table_data in (tables or []):
                    if not table_data:
                        continue
                    rows = ["<table>"]
                    for row in table_data:
                        cells = "".join(
                            f"<td>{_escape((cell or '').strip())}</td>"
                            for cell in (row or [])
                        )
                        rows.append(f"  <tr>{cells}</tr>")
                    rows.append("</table>")
                    parts.append("\n".join(rows))
                for t_obj in (page.find_tables() or []):
                    table_bboxes.append(t_obj.bbox)
            except Exception:
                pass

            try:
                chars = page.chars or []
                if chars:
                    avg_height = sum(c.get("height", 0) for c in chars) / len(chars)
                else:
                    avg_height = 0

                def _in_table(c) -> bool:
                    x0, y0, x1, y1 = c.get("x0", 0), c.get("top", 0), c.get("x1", 0), c.get("bottom", 0)
                    for tx0, ty0, tx1, ty1 in table_bboxes:
                        if x0 >= tx0 and x1 <= tx1 and y0 >= ty0 and y1 <= ty1:
                            return True
                    return False

                text_chars = [c for c in chars if not _in_table(c)]

                lines: dict[float, list] = {}
                for c in text_chars:
                    top = round(c.get("top", 0), 1)
                    lines.setdefault(top, []).append(c)

                for top_pos in sorted(lines):
                    line_chars = sorted(lines[top_pos], key=lambda c: c.get("x0", 0))
                    line_text = "".join(c.get("text", "") for c in line_chars).strip()
                    if not line_text:
                        continue
                    line_avg_h = sum(c.get("height", 0) for c in line_chars) / len(line_chars)
                    if avg_height > 0 and line_avg_h > avg_height * 1.2:
                        parts.append(f"<h2>{_escape(line_text)}</h2>")
                    else:
                        parts.append(f"<p>{_escape(line_text)}</p>")

            except Exception:
                raw = page.extract_text() or ""
                for line in raw.splitlines():
                    line = line.strip()
                    if line:
                        parts.append(f"<p>{_escape(line)}</p>")

    body_content = "\n".join(parts)
    return f"<html>\n<body>\n{body_content}\n</body>\n</html>\n"


# ---------------------------------------------------------------------------
# Field extraction (E-T1)
# ---------------------------------------------------------------------------

_FIELD_RE = re.compile(r"\{([a-zA-Z_][a-zA-Z0-9_.]*)\}")


def extract_fields(html: str, registry_path=None) -> list[dict]:
    """Extract {field_name} tokens from HTML (on plain text). Deduplicates.

    When registry_path is provided, dotted-path placeholders (e.g. {account.data.firstName})
    are resolved against the registry and their velocity path written into data_source.
    Unresolved dotted names get data_source = "UNRESOLVED:<name>" and a stderr warning.
    """
    try:
        from bs4 import BeautifulSoup
        text = BeautifulSoup(html, "html.parser").get_text()
    except ImportError:
        text = re.sub(r"<[^>]+>", " ", html)

    seen: dict[str, list[int]] = {}
    for i, m in enumerate(re.finditer(_FIELD_RE, text)):
        name = m.group(1)
        seen.setdefault(name, []).append(i)

    lookup: dict = {}
    if registry_path and any("." in name for name in seen):
        try:
            _scripts = str(Path(__file__).parent)
            if _scripts not in sys.path:
                sys.path.insert(0, _scripts)
            from agent_tools import build_velocity_lookup  # noqa: PLC0415
            lookup = build_velocity_lookup(registry_path)
        except Exception:
            pass

    unresolved: list[str] = []
    fields: list[dict] = []
    for name in seen:
        if "." in name and lookup:
            velocity = lookup.get(name)
            if velocity:
                data_source = velocity
            else:
                data_source = f"UNRESOLVED:{name}"
                unresolved.append(name)
        else:
            data_source = ""
        fields.append({
            "name": name,
            "token": f"$TBD_{name}",
            "data_source": data_source,
            "confidence": "",
        })

    if unresolved:
        print(
            f"WARNING: unresolved dotted placeholders: {', '.join(unresolved)}",
            file=sys.stderr,
        )

    return fields


def annotate_fields(html: str, fields: list[dict]) -> str:
    """Replace {field_name} → $TBD_field_name in the HTML string."""
    result = html
    for f in fields:
        result = result.replace("{" + f["name"] + "}", f["token"])
    return result


# ---------------------------------------------------------------------------
# Conditional block extraction (E-T2)
# ---------------------------------------------------------------------------

def _find_top_level_brackets(text: str) -> list[tuple[int, int, str]]:
    """Find top-level [[...]] blocks, skipping nested ones.
    Returns list of (start, end, inner_content) tuples.
    Handles cases like [[ outer [[inner]] still outer ]] correctly.
    """
    results = []
    i = 0
    n = len(text)
    while i < n - 1:
        if text[i : i + 2] == "[[":
            depth = 1
            start = i
            i += 2
            while i < n - 1 and depth > 0:
                if text[i : i + 2] == "[[":
                    depth += 1
                    i += 2
                elif text[i : i + 2] == "]]":
                    depth -= 1
                    if depth == 0:
                        results.append((start, i + 2, text[start + 2 : i]))
                    i += 2
                else:
                    i += 1
            if depth > 0:
                i = start + 2  # unclosed — skip
        else:
            i += 1
    return results


def extract_conditionals(html: str) -> list[dict]:
    """Extract [[conditional text]] blocks including nested ones.

    Nested [[...]] inside a block become child blocks with their own IDs.
    Parent source_text references children via $doc.condN.
    IDs are pre-order (parent lower than children); list is sorted at return.
    Each block carries top_level=True/False so annotate_conditionals can
    restrict HTML replacement to top-level blocks only.
    """
    try:
        from bs4 import BeautifulSoup
        text = BeautifulSoup(html, "html.parser").get_text()
    except ImportError:
        text = re.sub(r"<[^>]+>", " ", html)

    blocks: list[dict] = []
    next_id = [1]

    def _process_children(content: str) -> str:
        nested = _find_top_level_brackets(content)
        if not nested:
            return content.strip()
        result = content
        offset = 0
        for start, end, inner in nested:
            child_id = next_id[0]
            next_id[0] += 1
            child_source = _process_children(inner)
            blocks.append({"id": child_id, "source_text": child_source, "raw_text": inner, "top_level": False})
            ref = f"$doc.cond{child_id}"
            result = result[: start + offset] + ref + result[end + offset :]
            offset += len(ref) - (end - start)
        return result.strip()

    for _start, _end, content in _find_top_level_brackets(text):
        outer_id = next_id[0]
        next_id[0] += 1
        source = _process_children(content)
        blocks.append({"id": outer_id, "source_text": source, "raw_text": content, "top_level": True})

    blocks.sort(key=lambda b: b["id"])
    return blocks


def annotate_conditionals(html: str, blocks: list[dict]) -> str:
    """Replace [[...]] → [[...]]$doc.condN for all blocks (top-level and nested).

    Top-level blocks are matched by character position (right-to-left).
    Child blocks are matched by a naive string replace on their raw_text;
    this works as long as the child content contains no HTML tags.
    """
    regions = _find_top_level_brackets(html)
    top_level = [b for b in blocks if b.get("top_level", True)]

    if len(regions) != len(top_level):
        result = html
        for b in top_level:
            original = "[[" + b.get("raw_text", b["source_text"]) + "]]"
            result = result.replace(original, f"{original}$doc.cond{b['id']}")
    else:
        result = html
        for (start, end, content), b in zip(reversed(regions), reversed(top_level)):
            label = f"[[{content}]]"
            result = result[:start] + f"{label}$doc.cond{b['id']}" + result[end:]

    # Second pass: annotate nested child blocks within the already-annotated HTML.
    for b in [b for b in blocks if not b.get("top_level", True)]:
        raw = b.get("raw_text", b["source_text"])
        original = f"[[{raw}]]"
        result = result.replace(original, f"{original}$doc.cond{b['id']}")

    return result


def _normalise_for_leg2(fields: list[dict], source_name: str) -> dict:
    """Convert fields list to a leg2-compatible .mapping.yaml dict (uses placeholder)."""
    import datetime as dt

    variables = [
        {
            "name": f["name"],
            "placeholder": f["token"],
            "type": "variable",
            "context": {
                "parent_tag": "p",
                "line": None,
                "nearest_label": "",
            },
            "data_source": f.get("data_source", ""),
        }
        for f in fields
    ]
    return {
        "schema_version": "1.0",
        "source": source_name,
        "generated_at": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%S+00:00"),
        "variables": variables,
        "loops": [],
    }


def write_leg2_mapping(fields: list[dict], source_name: str, output_path: Path) -> None:
    """Write {stem}.mapping.yaml — leg2-compatible format with placeholder field."""
    data = _normalise_for_leg2(fields, source_name)
    output_path.write_text(
        yaml.dump(data, default_flow_style=False, allow_unicode=True),
        encoding="utf-8",
    )


# ---------------------------------------------------------------------------
# Write conditional form (E-T4)
# ---------------------------------------------------------------------------

_TBD_DISPLAY_RE = re.compile(r"\$TBD_([A-Za-z_][\w.]*)")


def _tbd_to_braces(text: str) -> str:
    """Display conversion: $TBD_name → {name} (customer-facing form only).

    Trailing dots are sentence punctuation, not part of the field name —
    `$TBD_amount.` becomes `{amount}.`, not `{amount.}`.
    """
    def _repl(m: re.Match) -> str:
        name = m.group(1)
        stripped = name.rstrip(".")
        return "{" + stripped + "}" + name[len(stripped):]
    return _TBD_DISPLAY_RE.sub(_repl, text)


def _braces_to_tbd(text: str) -> str:
    """Inverse of _tbd_to_braces: {name} → $TBD_name (canonical machine form).

    No-op on text already in $TBD_ form, so forms written before the {field}
    display change still parse.
    """
    return _FIELD_RE.sub(lambda m: "$TBD_" + m.group(1), text)


def write_conditional_form(blocks: list[dict], stem: str, output_path: Path) -> None:
    """Write {stem}.conditional-form.md — customer-facing conditional form."""
    lines = [
        f"# Conditional Text Review — {stem}",
        "",
        "For each block below, fill in the condition that controls when this text appears.",
        "Use accessor path format — dot notation without `$` or `()` syntax:",
        "",
        "| Root | Example accessor | Meaning |",
        "|------|-----------------|---------|",
        "| `quote` | `quote.quoteNumber` | Quote-level system fields |",
        "| `account` | `account.data.firstName` | Policyholder fields |",
        "| `policy` | `policy.data.riderType` | Custom policy fields |",
        "| `item` | `item.data.vin` | Per-exposure fields (within a loop) |",
        "",
        "Comparison examples: `quote.quoteNumber != null` · `account.data.state == \"CA\"` · `policy.data.riderType == \"AirBag\"`",
        "",
        "Run `python3 scripts/list_paths.py` to see all available accessors.",
        "Return this file to your implementation contact when complete.",
        "",
    ]
    for b in blocks:
        lines += [
            "---",
            "",
            f"## Block {b['id']}",
            "",
            f"> {_tbd_to_braces(b['source_text'])}",
            "",
            "Condition: ",
            "",
        ]
    output_path.write_text("\n".join(lines), encoding="utf-8")


# ---------------------------------------------------------------------------
# Parse filled conditional form → conditional-registry.yaml (E-T5)
# ---------------------------------------------------------------------------

def parse_conditional_form(md_path: Path) -> list[dict]:
    """Parse a customer-filled conditional form. Returns list of block dicts."""
    text = md_path.read_text(encoding="utf-8")
    blocks = []

    block_re = re.compile(
        r"##\s+Block\s+(\d+)\s*\n+>\s+(.+?)\s*\n+Condition:\s*([^\n]*)",
        re.DOTALL,
    )
    for m in block_re.finditer(text):
        block_id = int(m.group(1))
        source_text = _braces_to_tbd(m.group(2).strip())
        raw_condition = m.group(3).strip()
        # Take only the first line of the condition (customer may add notes below)
        condition_line = raw_condition.splitlines()[0].strip() if raw_condition else ""
        conditions = [condition_line] if condition_line else []
        blocks.append({
            "id": block_id,
            "source_text": source_text,
            "conditions": conditions,
            "operator": "AND",
        })

    return blocks


def write_conditional_registry(blocks: list[dict], output_path: Path) -> None:
    """Write {stem}.conditional-registry.yaml for Leg 4."""
    output_path.write_text(
        yaml.dump(blocks, default_flow_style=False, allow_unicode=True),
        encoding="utf-8",
    )


# ---------------------------------------------------------------------------
# Output writer
# ---------------------------------------------------------------------------

def write_output(html: str, input_path: Path, output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"{input_path.stem}.raw.html"
    output_path.write_text(html, encoding="utf-8")
    print(f"Wrote {output_path}")
    return output_path


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main() -> int:
    parser = argparse.ArgumentParser(
        description="Leg 0: convert PDF/Word to raw HTML and extract fields/conditionals."
    )
    parser.add_argument("--input", default=None, help="Path to .docx or .pdf file")
    parser.add_argument("--output-dir", default=None, help="Output directory (default: input's parent)")
    parser.add_argument(
        "--parse-conditional-form",
        default=None,
        metavar="FILLED_FORM.md",
        help="Parse a filled-in conditional form → conditional-registry.yaml",
    )
    args = parser.parse_args()

    # --- Mode: parse filled conditional form ---
    if args.parse_conditional_form:
        form_path = Path(args.parse_conditional_form)
        if not form_path.exists():
            print(f"Error: file not found: {form_path}", file=sys.stderr)
            return 1
        output_dir = Path(args.output_dir) if args.output_dir else form_path.parent
        output_dir.mkdir(parents=True, exist_ok=True)

        stem = form_path.stem
        if stem.endswith(".conditional-form"):
            stem = stem[: -len(".conditional-form")]

        blocks = parse_conditional_form(form_path)
        registry_path = output_dir / f"{stem}.conditional-registry.yaml"
        write_conditional_registry(blocks, registry_path)
        print(f"Wrote {registry_path}")
        return 0

    # --- Mode: convert document ---
    if not args.input:
        print("Error: --input is required (unless using --parse-conditional-form)", file=sys.stderr)
        return 1

    input_path = Path(args.input)

    if not input_path.exists():
        print(f"Error: file not found: {input_path}", file=sys.stderr)
        return 1

    suffix = input_path.suffix.lower()

    if suffix == ".doc":
        print(
            "Error: legacy .doc format is not supported.\n"
            "Open the file in Microsoft Word → File → Save As → .docx, then re-run.",
            file=sys.stderr,
        )
        return 1

    if suffix not in (".docx", ".pdf"):
        print(
            f"Error: unsupported format '{suffix}'. Accepted: .docx, .pdf",
            file=sys.stderr,
        )
        return 1

    output_dir = Path(args.output_dir) if args.output_dir else input_path.parent
    output_dir.mkdir(parents=True, exist_ok=True)
    stem = input_path.stem

    # Find registry for explicit path resolution (walk up from input dir)
    registry_path = None
    cur = input_path.parent.resolve()
    for _ in range(8):
        for rel in ("registry/path-registry.yaml", "path-registry.yaml"):
            candidate = cur / rel
            if candidate.is_file():
                registry_path = candidate
                break
        if registry_path:
            break
        if cur.parent == cur:
            break
        cur = cur.parent

    # Convert to raw HTML
    if suffix == ".docx":
        raw_html = convert_docx(input_path)
    else:
        raw_html = convert_pdf(input_path)

    # Write raw HTML
    raw_path = output_dir / f"{stem}.raw.html"
    raw_path.write_text(raw_html, encoding="utf-8")
    print(f"Wrote {raw_path}")

    # Extract + annotate fields
    fields = extract_fields(raw_html, registry_path=registry_path)
    annotated = annotate_fields(raw_html, fields)

    # Extract + annotate conditionals
    blocks = extract_conditionals(annotated)
    annotated = annotate_conditionals(annotated, blocks)

    # Write annotated HTML (pipeline input for Leg 1 / Leg 3)
    annotated_path = output_dir / f"{stem}.annotated.html"
    annotated_path.write_text(annotated, encoding="utf-8")
    print(f"Wrote {annotated_path}")

    # Write leg2-compatible mapping
    mapping_path = output_dir / f"{stem}.mapping.yaml"
    write_leg2_mapping(fields, f"{stem}.annotated.html", mapping_path)
    print(f"Wrote {mapping_path}")

    # Write conditional form (only if there are conditionals)
    if blocks:
        form_path = output_dir / f"{stem}.conditional-form.md"
        write_conditional_form(blocks, stem, form_path)
        print(f"Wrote {form_path}")
    else:
        print("No [[conditional]] blocks found — skipping conditional form.")

    return 0


if __name__ == "__main__":
    sys.exit(main())
